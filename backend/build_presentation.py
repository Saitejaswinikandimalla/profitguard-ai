"""
build_presentation.py
Populates the official DATATHON-Round1-Idea-PPT-Template.pptx with ProfitGuard AI details,
metrics, and embeds high-resolution prototype screenshots directly into the slides.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

TEMPLATE_PATH = r"C:\Users\tejas\Downloads\DATATHON-Round1-Idea-PPT-Template.pptx"
OUTPUT_PATH = r"C:\Users\tejas\Downloads\ProfitGuard_AI_DATATHON_Round1.pptx"
IMG_DIR = r"C:\Users\tejas\.gemini\antigravity\scratch\profitguard_work\screenshots"


def fill_presentation():
    if not os.path.exists(TEMPLATE_PATH):
        print(f"Template not found at {TEMPLATE_PATH}")
        return

    prs = Presentation(TEMPLATE_PATH)

    # -------------------------------------------------------------
    # SLIDE 1: Title & Team Details
    # -------------------------------------------------------------
    s1 = prs.slides[0]
    for shape in s1.shapes:
        if shape.name == "TextBox 7":
            tf = shape.text_frame
            tf.clear()

            details = [
                ("Team Name", "ProfitGuard AI Team"),
                ("Team Members", "Tejas & Collaborator (Track 4 Team)"),
                ("Track Chosen", "Track 4: Business — The Money Leak"),
                ("Problem Statement", "Revenue Leakage Detection & Decision Support"),
                ("Product Name", "ProfitGuard AI")
            ]

            for i, (label, val) in enumerate(details):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"{label}: "
                p.font.name = "Arial"
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = RGBColor(15, 23, 42)

                run = p.add_run()
                run.text = val
                run.font.name = "Arial"
                run.font.size = Pt(20)
                run.font.bold = False
                run.font.color.rgb = RGBColor(30, 41, 59)
                p.space_after = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 2: Problem & Target User
    # -------------------------------------------------------------
    s2 = prs.slides[1]
    for shape in s2.shapes:
        if shape.name == "TextBox 8":
            tf = shape.text_frame
            tf.clear()

            points = [
                ("Target User & Operational Pain Point",
                 "Multichannel retail and e-commerce managers lack unified visibility across transaction pricing, vendor costs, and customer behavior. Critical financial leaks slip past unnoticed."),
                ("Why This Problem Matters",
                 "Uncontrolled discounting, catalog pricing errors, and high product return rates directly erode thin retail operating margins, converting apparent revenue growth into hidden cash losses."),
                ("Hard Dataset Evidence (Datathon Business Track)",
                 "Analysis of 18,000 transactions reveals 5,315 completed sales (38.07%) operated at negative estimated margins, causing an alarming ₹12.14 Lakhs in cumulative revenue leakage out of ₹73.02 Lakhs total sales.")
            ]

            for i, (heading, body) in enumerate(points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {heading}: "
                p.font.name = "Arial"
                p.font.size = Pt(16)
                p.font.bold = True
                p.font.color.rgb = RGBColor(15, 23, 42)

                run = p.add_run()
                run.text = body
                run.font.name = "Arial"
                run.font.size = Pt(15)
                run.font.bold = False
                run.font.color.rgb = RGBColor(51, 65, 85)
                p.space_after = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 3: Solution & USP
    # -------------------------------------------------------------
    s3 = prs.slides[2]
    for shape in s3.shapes:
        if shape.name == "TextBox 8":
            tf = shape.text_frame
            tf.clear()
            sol_points = [
                ("Unified Financial Risk Platform",
                 "ProfitGuard AI bridges transaction sales, product costs, inventory levels, and campaign records into a single decision-support intelligence hub."),
                ("Transparent Unit Economics Auditing",
                 "Instantly evaluates unit-level economics on every transaction: Q × [P(1 - D) - C], automatically isolating catalog mispricings vs. promo discount leakages.")
            ]
            for i, (h, b) in enumerate(sol_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}: "
                p.font.name = "Arial"
                p.font.size = Pt(15)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(14)
                run.font.bold = False
                p.space_after = Pt(8)

        elif shape.name == "TextBox 10":
            tf = shape.text_frame
            tf.clear()
            usp_points = [
                ("Deterministic Math + From-Scratch ML",
                 "Unlike black-box AI tools, ProfitGuard uses fully explainable audit logic for revenue leakage paired with an independent NumPy K-Means model for customer segmentation."),
                ("Root-Cause Prescription vs Passive Reporting",
                 "Differentiates between catalog master-data errors (Price < Cost) and promotional leakages (Discount pushing below cost), prescribing actionable fixes for operations.")
            ]
            for i, (h, b) in enumerate(usp_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}: "
                p.font.name = "Arial"
                p.font.size = Pt(15)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(14)
                run.font.bold = False
                p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 4: Dataset & Insights (With Executive Chart Image)
    # -------------------------------------------------------------
    s4 = prs.slides[3]
    for shape in s4.shapes:
        if shape.name == "TextBox 8":
            shape.width = Inches(6.0)
            shape.height = Inches(4.3)
            tf = shape.text_frame
            tf.clear()
            dataset_points = [
                ("Multi-Table Relational Schema",
                 "Joined transactions.csv (18k rows) with products.csv (250 rows) on product_id; integrated customers.csv (3k), inventory.csv (5k), and campaigns.csv (1.2k)."),
                ("Methodological Assumptions",
                 "Filtered status == 'Completed' (13,962 rows). Assumed unit_price represents list/pre-discount price; effective selling price = P × (1 - D)."),
                ("Critical Empirical Discoveries",
                 "• 4,770 transactions had Catalog Mispricing (unit_price < unit_cost before discount).\n• 545 transactions suffered Discount Leakage (margin turned negative by steep promo).\n• Completed Revenue: ₹73.02 Lakhs | Total Leakage: ₹12.14 Lakhs (16.6% lost!).")
            ]
            for i, (h, b) in enumerate(dataset_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}: "
                p.font.name = "Arial"
                p.font.size = Pt(13.5)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(12.5)
                run.font.bold = False
                p.space_after = Pt(8)

    img1_path = os.path.join(IMG_DIR, "screen1_executive.png")
    if os.path.exists(img1_path):
        s4.shapes.add_picture(img1_path, Inches(6.9), Inches(2.2), Inches(5.9), Inches(3.8))

    # -------------------------------------------------------------
    # SLIDE 5: ML Model (With K-Means Cluster Image)
    # -------------------------------------------------------------
    s5 = prs.slides[4]
    for shape in s5.shapes:
        if shape.name == "TextBox 8":
            shape.width = Inches(6.0)
            shape.height = Inches(4.3)
            tf = shape.text_frame
            tf.clear()
            ml_points = [
                ("Model Built From Scratch (NumPy)",
                 "Custom vectorized K-Means implementation featuring random multi-restart (n_init=5), Euclidean assignment, and convergence tolerance Δ < 1e-4."),
                ("RFM Feature Matrix & Normalization",
                 "Extracted Recency (days), Frequency (orders), and Monetary (spend) across 2,971 purchasing customers, scaled via from-scratch Z-Score Standardization."),
                ("Baseline & Metric Evaluation",
                 "Evaluated against static heuristic rules using Within-Cluster Sum of Squares (Inertia) & sample-based Silhouette Score (s ≈ 0.38)."),
                ("Product Utilization & Business Action",
                 "Connects behavioral clusters directly to pricing policies: protects High-Value Champions while enforcing strict discount caps on deal seekers.")
            ]
            for i, (h, b) in enumerate(ml_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}: "
                p.font.name = "Arial"
                p.font.size = Pt(13.5)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(12.5)
                run.font.bold = False
                p.space_after = Pt(8)

    img3_path = os.path.join(IMG_DIR, "screen3_ml_clusters.png")
    if os.path.exists(img3_path):
        s5.shapes.add_picture(img3_path, Inches(6.9), Inches(2.2), Inches(5.9), Inches(3.8))

    # -------------------------------------------------------------
    # SLIDE 6: Base Prototype & MVP (With Leakage Audit Image)
    # -------------------------------------------------------------
    s6 = prs.slides[5]
    for shape in s6.shapes:
        if shape.name == "TextBox 8":
            shape.width = Inches(6.0)
            tf = shape.text_frame
            tf.clear()
            proto_points = [
                ("Working Interactive Prototype (Built So Far)",
                 "• Executive Overview: Real-time macro KPIs, category margins, and Top 10 worst leaks.\n• Revenue Leakage Detector: Transaction auditor with step-by-step mathematical breakdown.\n• Customer Intelligence: Interactive NumPy K-Means clustering with 2D/3D Plotly visualizer.")
            ]
            for i, (h, b) in enumerate(proto_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}:\n"
                p.font.name = "Arial"
                p.font.size = Pt(13)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(12)
                run.font.bold = False
                p.space_after = Pt(6)

        elif shape.name == "TextBox 10":
            shape.width = Inches(6.0)
            tf = shape.text_frame
            tf.clear()
            mvp_points = [
                ("Round 2 Must-Have MVP Roadmap",
                 "1. Automated Webhook Alerts: Instant Slack/Email notification when transaction margin < 0.\n2. Churn Prediction Model: Supervised classification model predicting customer lapse risk.\n3. Dynamic Discount Guardrails: API middleware enforcing profit-floor limits at checkout."),
                ("End-to-End Decision Workflow",
                 "Input: Raw ERP/POS feeds → Processing: Margin audit & RFM clustering → Action: Auto-flagged priority investigation queue with operational fix recommendations.")
            ]
            for i, (h, b) in enumerate(mvp_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}:\n"
                p.font.name = "Arial"
                p.font.size = Pt(13)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(12)
                run.font.bold = False
                p.space_after = Pt(6)

    img2_path = os.path.join(IMG_DIR, "screen2_leakage.png")
    if os.path.exists(img2_path):
        s6.shapes.add_picture(img2_path, Inches(6.9), Inches(2.2), Inches(5.9), Inches(3.6))

    # -------------------------------------------------------------
    # SLIDE 7: Technical Approach
    # -------------------------------------------------------------
    s7 = prs.slides[6]
    for shape in s7.shapes:
        if shape.name == "TextBox 8":
            tf = shape.text_frame
            tf.clear()
            tech_points = [
                ("Core Technologies & Frameworks",
                 "• Core Language: Python 3.14\n• Analytics & Relational Processing: Pandas, NumPy\n• Custom ML Algorithm: Pure NumPy (K-Means, Euclidean Distance, Z-Score Scaler)\n• Interactive Web Framework: Streamlit\n• Data Visualization: Plotly Express & Plotly Graph Objects"),
                ("Architecture & End-to-End Pipeline",
                 "1. Data Layer: Ingests 5 CSV datasets (Transactions, Products, Customers, Inventory, Campaigns).\n2. ETL & Auditing Engine: Joins transactions to cost tables and executes unit economics formulas.\n3. ML Intelligence: Vectorized RFM feature extraction and unsupervised behavioral clustering.\n4. Presentation & Decision Layer: Real-time interactive Streamlit dashboard for business stakeholders.")
            ]
            for i, (h, b) in enumerate(tech_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}:\n"
                p.font.name = "Arial"
                p.font.size = Pt(15)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(14)
                run.font.bold = False
                p.space_after = Pt(12)

    # -------------------------------------------------------------
    # SLIDE 8: Feasibility, Impact & Next Steps
    # -------------------------------------------------------------
    s8 = prs.slides[7]
    for shape in s8.shapes:
        if shape.name == "TextBox 8":
            tf = shape.text_frame
            tf.clear()
            final_points = [
                ("Feasibility & Operational Risk Mitigation",
                 "• Feasibility: Zero external paid API or GPU requirements; runs efficiently in lightweight CPU environments.\n• Risk: Real-world transactions have tax, shipping, and returns complexities → Mitigated via modular schema adapters."),
                ("Impact & Synthetic Data Limitations",
                 "• Business Impact: Recovers 5–15% of bottom-line net margin by catching pricing leaks before order dispatch.\n• Data Limitation: Synthetic educational dataset contains decoupled price/cost distributions; production deployment will calibrate against actual ERP cost accounting."),
                ("Next Steps & Academic References",
                 "• Round 2 Roadmap: PostgreSQL backend integration, real-time alerting webhooks, supervised churn classifier.\n• References: RFM Customer Analysis (Hughes, 1994); K-Means Clustering (MacQueen, 1967); Oracle Retail Analytics.")
            ]
            for i, (h, b) in enumerate(final_points):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {h}: "
                p.font.name = "Arial"
                p.font.size = Pt(15)
                p.font.bold = True
                run = p.add_run()
                run.text = b
                run.font.name = "Arial"
                run.font.size = Pt(14)
                run.font.bold = False
                p.space_after = Pt(10)

    prs.save(OUTPUT_PATH)
    print(f"Presentation successfully updated with embedded screenshots at: {OUTPUT_PATH}")


if __name__ == "__main__":
    fill_presentation()
